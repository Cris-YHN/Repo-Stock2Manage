import sqlite3
from openpyxl import load_workbook

DB_PATH = "database/s2m.db"
XLSX_PATH = "assets/Excels/codigos_postales_Argentina_Unicos.xlsx"

def cargar_codigos_postales():
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()

    wb = load_workbook(XLSX_PATH, read_only=True, data_only=True)
    sheet = wb.active  # primera hoja

    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not row or all(c is None for c in row):
            continue  # salta filas vacías

        # Obtener valores y asegurar longitud
        valores = list(row) + [None] * (3 - len(row))
        codigo, ciudad, provincia = valores[:3]
        pais = "Argentina"

        # 🔧 Normalizar código postal
        if codigo is None or ciudad is None:
            print(f"⚠️ Fila incompleta, se omite: {valores}")
            continue

        if isinstance(codigo, float):
            codigo = str(int(codigo))
        else:
            codigo = str(codigo).strip()

        try:
            cursor.execute(
                """
                INSERT INTO codigos_postales (codigo_postal, ciudad, provincia, pais)
                VALUES (?, ?, ?, ?)
                """,
                (codigo, ciudad, provincia, pais),
            )
        except sqlite3.IntegrityError:
            print(f"⚠️ Código {codigo} ya existe, se salta")

    conn.commit()
    conn.close()
    print("✅ Códigos postales cargados correctamente")

if __name__ == "__main__":
    cargar_codigos_postales()
