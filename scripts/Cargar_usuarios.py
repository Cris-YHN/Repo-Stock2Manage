import sqlite3
import bcrypt
from openpyxl import load_workbook

DB_PATH = "database/s2m.db"
XLSX_PATH = "assets/Excels/usuarios_import.xlsx"

def cargar_usuarios():
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()

    wb = load_workbook(XLSX_PATH, read_only=True, data_only=True)
    sheet = wb.active

    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not row or all(c is None for c in row):
            continue

        nombre, apellido, email, puesto, activo = row

        # Normalizar el puesto a minúsculas
        puesto = puesto.lower().strip()

        # Generar contraseña temporal (1234)
        contrasenia = hash_password("1234")

        try:
            cursor.execute("""
                INSERT INTO usuarios (nombre, apellido, contrasenia, email, puesto, activo)
                VALUES (?, ?, ?, ?, ?, ?)
            """, (nombre, apellido, contrasenia, email, puesto, activo))
        except sqlite3.IntegrityError:
            print(f"⚠️ Usuario {email} ya existe, se salta")

    conn.commit()
    conn.close()
    print("✅ Usuarios cargados correctamente")

def hash_password(password: str) -> str:
        # Genera el hash con salt
        salt = bcrypt.gensalt()
        hashed = bcrypt.hashpw(password.encode("utf-8"), salt)
        return hashed.decode("utf-8")  # se guarda como texto en la DB

if __name__ == "__main__":
    cargar_usuarios()
